import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6] # blank layout
    
    # Colors
    BG_COLOR = RGBColor(9, 13, 22)         # #090d16
    CARD_BG = RGBColor(17, 24, 39)          # #111827
    CARD_BORDER = RGBColor(30, 41, 59)      # #1e293b
    CYAN = RGBColor(6, 182, 212)            # #06b6d4
    WHITE = RGBColor(248, 250, 252)         # #f8fafc
    MUTED = RGBColor(148, 163, 184)         # #94a3b8
    GREEN = RGBColor(34, 197, 94)           # #22c55e
    AMBER = RGBColor(245, 158, 11)          # #f59e0b
    RED = RGBColor(239, 68, 68)             # #ef4444

    def set_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="AQUASENTINEL"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN
        p.font.name = 'Arial'

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.font.name = 'Arial'

    # ==========================================
    # SLIDE 1: Title & Hero Hook
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide1)
    
    # Title Box
    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "AquaSentinel"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.font.name = 'Arial'
    
    p2 = tf1.add_paragraph()
    p2.text = "INTELLIGENT WATER OPERATIONS PLATFORM"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = MUTED
    p2.space_before = Pt(10)
    
    p3 = tf1.add_paragraph()
    p3.text = '"AI that detects water waste before it becomes a bill."'
    p3.font.size = Pt(24)
    p3.font.italic = True
    p3.font.color.rgb = WHITE
    p3.space_before = Pt(24)
    
    p4 = tf1.add_paragraph()
    p4.text = "An intelligent water operations platform that monitors telemetry, detects abnormal usage, explains the signal, predicts potential impact, and helps facility teams act."
    p4.font.size = Pt(15)
    p4.font.color.rgb = MUTED
    p4.space_before = Pt(20)

    # Demo Tag shape
    shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(6.0), Inches(5.5), Inches(0.6))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CYAN
    tf_tag = shape.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "DEMO ENVIRONMENT — SIMULATED SENSOR TELEMETRY"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = CYAN
    p_tag.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 2: The Invisible Problem
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide2)
    add_header(slide2, "Water Waste Is Often Invisible Until It Becomes Expensive", "PROBLEM STATEMENT")

    # Subtitle
    sub_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = '"Water leaks and abnormal consumption can continue unnoticed, turning a small deviation into significant water and financial loss."'
    p_sub.font.size = Pt(15)
    p_sub.font.italic = True
    p_sub.font.color.rgb = MUTED

    # 4 Cards Layout
    card_w, card_h = Inches(5.6), Inches(1.8)
    coords = [
        (Inches(0.8), Inches(2.1)), (Inches(6.9), Inches(2.1)),
        (Inches(0.8), Inches(4.1)), (Inches(6.9), Inches(4.1))
    ]
    challenges = [
        ("1. Hidden Leaks", "Sub-surface and behind-the-wall pipe breaches continue without immediate visibility."),
        ("2. Valve Overflows", "Tank float and hydrostatic valve failures cause continuous reservoir overspills."),
        ("3. Off-Peak Flow", "Night-time continuous running fixtures are difficult to identify manually."),
        ("4. Context Gap", "Facility teams often receive raw meter readings without actionable operational context.")
    ]
    for (x, y), (ctitle, cdesc) in zip(coords, challenges):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_ct = tf_c.paragraphs[0]
        p_ct.text = ctitle
        p_ct.font.size = Pt(16)
        p_ct.font.bold = True
        p_ct.font.color.rgb = CYAN
        
        p_cd = tf_c.add_paragraph()
        p_cd.text = cdesc
        p_cd.font.size = Pt(13)
        p_cd.font.color.rgb = WHITE
        p_cd.space_before = Pt(8)

    # Key Insight Banner
    banner = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.8))
    banner.fill.solid()
    banner.fill.fore_color.rgb = CARD_BG
    banner.line.color.rgb = AMBER
    tf_b = banner.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "Raw Telemetry Data  ≠  Actionable Operational Decisions"
    p_b.font.size = Pt(18)
    p_b.font.bold = True
    p_b.font.color.rgb = AMBER
    p_b.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 3: Solution
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide3)
    add_header(slide3, "Closing the Gap From Raw Signal to Maintenance Decision", "OUR SOLUTION")

    steps = [
        ("MONITOR", "Continuous Telemetry Visibility", "1.0s real-time stream of flow, pressure & level"),
        ("DETECT", "Deterministic Baseline Engine", "Flags sustained flow spikes > 1.5x baseline"),
        ("EXPLAIN", "AI Diagnostic Reasoning", "Generates diagnostic cause & evidence timeline"),
        ("PREDICT", "Projected Impact (₹)", "Calculates water volume & INR financial loss"),
        ("ACT", "Integrated Work Orders", "Dispatches maintenance teams directly")
    ]
    step_w = Inches(2.2)
    gap = Inches(0.18)
    start_x = Inches(0.8)

    for i, (stitle, ssub, sdesc) in enumerate(steps):
        sx = start_x + i * (step_w + gap)
        scard = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.8), step_w, Inches(3.6))
        scard.fill.solid()
        scard.fill.fore_color.rgb = CARD_BG
        scard.line.color.rgb = CYAN if i in [1, 2] else CARD_BORDER
        tf_s = scard.text_frame
        tf_s.word_wrap = True
        
        p_st = tf_s.paragraphs[0]
        p_st.text = f"0{i+1}. {stitle}"
        p_st.font.size = Pt(16)
        p_st.font.bold = True
        p_st.font.color.rgb = CYAN
        
        p_ss = tf_s.add_paragraph()
        p_ss.text = ssub
        p_ss.font.size = Pt(13)
        p_ss.font.bold = True
        p_ss.font.color.rgb = WHITE
        p_ss.space_before = Pt(10)

        p_sd = tf_s.add_paragraph()
        p_sd.text = sdesc
        p_sd.font.size = Pt(11)
        p_sd.font.color.rgb = MUTED
        p_sd.space_before = Pt(10)

    # Core relationship box
    rel_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.1))
    rel_box.fill.solid()
    rel_box.fill.fore_color.rgb = CARD_BG
    rel_box.line.color.rgb = GREEN
    tf_r = rel_box.text_frame
    tf_r.word_wrap = True
    p_r1 = tf_r.paragraphs[0]
    p_r1.text = "THE CORE TECHNICAL STORY:"
    p_r1.font.size = Pt(11)
    p_r1.font.bold = True
    p_r1.font.color.rgb = GREEN
    
    p_r2 = tf_r.add_paragraph()
    p_r2.text = '"The deterministic anomaly engine identifies abnormal flow. AI then explains the detected anomaly and recommends an action."'
    p_r2.font.size = Pt(16)
    p_r2.font.bold = True
    p_r2.font.color.rgb = WHITE
    p_r2.space_before = Pt(4)

    # ==========================================
    # SLIDE 4: Architecture
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide4)
    add_header(slide4, "End-to-End System Architecture", "SYSTEM ARCHITECTURE")

    # Left box: Production Oriented
    pbox = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    pbox.fill.solid()
    pbox.fill.fore_color.rgb = CARD_BG
    pbox.line.color.rgb = CYAN
    tf_p = pbox.text_frame
    tf_p.word_wrap = True
    pp1 = tf_p.paragraphs[0]
    pp1.text = "PRODUCTION-ORIENTED ARCHITECTURE"
    pp1.font.size = Pt(16)
    pp1.font.bold = True
    pp1.font.color.rgb = CYAN
    
    p_items = [
        ("Flow Sensors / ESP32", "Physical pulse-based flow meters and hardware nodes"),
        ("MQTT / Modbus Protocol", "Industrial IoT messaging protocol layer"),
        ("Telemetry Stream Layer", "1.0s sampling real-time ingestion queue"),
        ("Deterministic Engine", "Threshold comparison (>1.5x baseline)"),
        ("AI Diagnostic Layer", "Root-cause explanation & evidence timeline"),
        ("Operations Dashboard", "Alerts & integrated work order dispatch")
    ]
    for title, desc in p_items:
        p_t = tf_p.add_paragraph()
        p_t.text = f"• {title}: {desc}"
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(10)

    # Right box: Current Prototype
    cbox = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
    cbox.fill.solid()
    cbox.fill.fore_color.rgb = CARD_BG
    cbox.line.color.rgb = AMBER
    tf_c = cbox.text_frame
    tf_c.word_wrap = True
    cp1 = tf_c.paragraphs[0]
    cp1.text = "CURRENT HACKATHON PROTOTYPE"
    cp1.font.size = Pt(16)
    cp1.font.bold = True
    cp1.font.color.rgb = AMBER

    c_items = [
        ("Simulated Telemetry Engine", "Powered by a standalone simulated telemetry engine to demonstrate the complete end-to-end intelligence workflow."),
        ("Deterministic Logic", "Strict 43 L/min baseline and 64.5 L/min anomaly threshold rules."),
        ("AI Diagnostic Engine", "Downstream diagnostic explanations powered by contextual reasoning."),
        ("Financial Loss Engine", "Real-time mathematical forecasting in Indian Rupees (₹)."),
        ("Work Order Workflow", "Full ticketing lifecycle from creation to resolution.")
    ]
    for title, desc in c_items:
        cp_t = tf_c.add_paragraph()
        cp_t.text = f"• {title}"
        cp_t.font.size = Pt(13)
        cp_t.font.bold = True
        cp_t.font.color.rgb = WHITE
        cp_t.space_before = Pt(12)
        
        cp_d = tf_c.add_paragraph()
        cp_d.text = desc
        cp_d.font.size = Pt(11)
        cp_d.font.color.rgb = MUTED
        cp_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 5: Technical Engine
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide5)
    add_header(slide5, "Detection Is Deterministic. AI Explains It.", "DETECTION ENGINE")

    # Mathematical Baseline Cards
    bcard1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(1.4))
    bcard1.fill.solid()
    bcard1.fill.fore_color.rgb = CARD_BG
    bcard1.line.color.rgb = CARD_BORDER
    tf_b1 = bcard1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "EXPECTED BASELINE (B)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = MUTED
    p2 = tf_b1.add_paragraph()
    p2.text = "43 L/min"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = CYAN
    p2.space_before = Pt(4)

    bcard2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(1.4))
    bcard2.fill.solid()
    bcard2.fill.fore_color.rgb = CARD_BG
    bcard2.line.color.rgb = CARD_BORDER
    tf_b2 = bcard2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "ANOMALY THRESHOLD (T = 1.5 × B)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = MUTED
    p2 = tf_b2.add_paragraph()
    p2.text = "64.5 L/min"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = RED
    p2.space_before = Pt(4)

    # Flow Progression Card
    pcard = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.3), Inches(11.7), Inches(2.6))
    pcard.fill.solid()
    pcard.fill.fore_color.rgb = CARD_BG
    pcard.line.color.rgb = CARD_BORDER
    tf_prog = pcard.text_frame
    tf_prog.word_wrap = True
    
    pp = tf_prog.paragraphs[0]
    pp.text = "REAL-TIME FLOW PROGRESSION STATE"
    pp.font.size = Pt(14)
    pp.font.bold = True
    pp.font.color.rgb = WHITE
    
    stages = [
        ("43 L/min", "Normal Baseline", GREEN),
        ("61 L/min", "Elevated Flow Variance", AMBER),
        ("78 L/min", "Threshold Crossed (>64.5 L/m)", RED),
        ("94 L/min", "Sustained Critical Anomaly", RED)
    ]
    for flow, desc, color in stages:
        pt = tf_prog.add_paragraph()
        pt.text = f"• {flow}  ───>  {desc}"
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = color
        pt.space_before = Pt(10)

    # Principle Box
    prin = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    prin.fill.solid()
    prin.fill.fore_color.rgb = CARD_BG
    prin.line.color.rgb = CYAN
    tf_pr = prin.text_frame
    tf_pr.word_wrap = True
    ppr = tf_pr.paragraphs[0]
    ppr.text = 'TECHNICAL CREDIBILITY PRINCIPLE:'
    ppr.font.size = Pt(11)
    ppr.font.bold = True
    ppr.font.color.rgb = CYAN
    ppr2 = tf_pr.add_paragraph()
    ppr2.text = '"AI is not guessing whether an anomaly exists. The deterministic engine compares live flow against a defined baseline threshold and evaluates sustained deviation; AI explains what the signal means."'
    ppr2.font.size = Pt(13)
    ppr2.font.bold = True
    ppr2.font.color.rgb = WHITE
    ppr2.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: Live Demo Comparison
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide6)
    add_header(slide6, "Normal vs Simulated Leak State (Prototype Telemetry)", "LIVE DEMO STATE")

    # Table creation
    rows, cols = 11, 3
    table_shape = slide6.shapes.add_table(rows, cols, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
    table = table_shape.table
    table.columns[0].width = Inches(3.7)
    table.columns[1].width = Inches(4.0)
    table.columns[2].width = Inches(4.0)

    headers = ["Parameter", "Normal System State", "Simulated Leak Scenario"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BORDER
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = CYAN

    data_rows = [
        ("Current Flow Rate", "43 L/min (Normal noise 41–44 L/min)", "94 L/min (Elevated ~2.2x baseline)"),
        ("Expected Baseline", "43 L/min", "43 L/min"),
        ("Deviation %", "0%", "+118%"),
        ("Threshold (1.5x)", "64.5 L/min", "64.5 L/min"),
        ("System Status", "Normal State", "Critical Abnormal-Flow State"),
        ("Active Alarms", "0 Alarms", "1 Active Alarm (Block B)"),
        ("15 min Simulated Excess", "0 L", "765 L"),
        ("6-Hour Projected Excess", "0 L", "18,360 L"),
        ("6-Hour Financial Impact", "₹0.00", "₹826.20 (at ₹45.00 / kL)"),
        ("Maintenance Action", "No action required", "Work Order #104 Triggered")
    ]
    for i, row in enumerate(data_rows):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.color.rgb = RED if (j==2 and i in [0,2,4,5,7,8,9]) else WHITE
            if j == 0:
                p.font.bold = True

    # ==========================================
    # SLIDE 7: AI Reasoning
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide7)
    add_header(slide7, "Don't Just Alert. Explain.", "AI REASONING ENGINE")

    # Diagnostic Summary Box
    diag_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.4))
    diag_box.fill.solid()
    diag_box.fill.fore_color.rgb = CARD_BG
    diag_box.line.color.rgb = CYAN
    tf_d = diag_box.text_frame
    tf_d.word_wrap = True

    p = tf_d.paragraphs[0]
    p.text = "OPERATIONAL DIAGNOSTIC SUMMARY  │  AI Confidence: 94%"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p = tf_d.add_paragraph()
    p.text = 'DIAGNOSTIC EXPLANATION:\n"Water flow has remained sustained at 94 L/min (2.2× baseline) for 15 min (simulated). The sustained continuous flow pattern is inconsistent with standard usage and indicates a continuous pipe leak."'
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.space_before = Pt(10)

    p = tf_d.add_paragraph()
    p.text = "EVIDENCE TIMELINE (MAGNITUDE + DURATION):\n• NORMAL: 43 L/min baseline\n• FLOW INCREASE: 61 L/min (Elevated consumption)\n• THRESHOLD CROSSED: 78 L/min (Exceeded 64.5 L/min)\n• SUSTAINED DEVIATION: 94 L/min (15 min simulated)\n• ANOMALY CONFIRMED: Sustained abnormal flow confirmed"
    p.font.size = Pt(12)
    p.font.color.rgb = MUTED
    p.space_before = Pt(10)

    p = tf_d.add_paragraph()
    p.text = "POSSIBLE CAUSE: Continuous pipe leakage\nRECOMMENDED ACTION: Inspect Block B — Floor 2 washroom & isolate valve"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p.space_before = Pt(10)

    # Comparison Banner
    banner7 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    banner7.fill.solid()
    banner7.fill.fore_color.rgb = CARD_BG
    banner7.line.color.rgb = CARD_BORDER
    tf_b7 = banner7.text_frame
    p_b7 = tf_b7.paragraphs[0]
    p_b7.text = "CONVENTIONAL: \"Flow is high.\"  ──>  AQUASENTINEL: \"Flow is high ──> Baseline exceeded ──> Deviation sustained ──> Possible cause explained ──> Impact projected ──> Action recommended.\""
    p_b7.font.size = Pt(12)
    p_b7.font.bold = True
    p_b7.font.color.rgb = WHITE
    p_b7.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 8: Actionable Response
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide8)
    add_header(slide8, "From Detection to Work Order Dispatch", "ACTIONABLE RESPONSE")

    # Ticket Card
    tcard = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.4))
    tcard.fill.solid()
    tcard.fill.fore_color.rgb = CARD_BG
    tcard.line.color.rgb = RED
    tf_t = tcard.text_frame
    tf_t.word_wrap = True

    p = tf_t.paragraphs[0]
    p.text = "WORK ORDER TICKET #104  │  PRIORITY: CRITICAL  │  STATUS: ASSIGNED"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RED

    t_details = [
        ("Facility Location", "Block B — Floor 2 Washrooms & Utility Line"),
        ("Telemetry Readout", "94 L/min current flow vs 43 L/min baseline (+118% deviation)"),
        ("Simulated Duration", "15 min simulated"),
        ("Projected 6-Hour Excess", "18,360 Liters"),
        ("Projected Financial Impact", "₹826.20 (configured at ₹45.00 / kL)"),
        ("Assigned Team", "Facility Maintenance / Plumbing Team")
    ]
    for label, val in t_details:
        pt = tf_t.add_paragraph()
        pt.text = f"• {label}: {val}"
        pt.font.size = Pt(14)
        pt.font.color.rgb = WHITE
        pt.space_before = Pt(10)

    # Reset Note
    rbox = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    rbox.fill.solid()
    rbox.fill.fore_color.rgb = CARD_BG
    rbox.line.color.rgb = CYAN
    tf_r = rbox.text_frame
    pr = tf_r.paragraphs[0]
    pr.text = 'SYSTEM RESET BEHAVIOR:'
    pr.font.size = Pt(11)
    pr.font.bold = True
    pr.font.color.rgb = CYAN
    pr2 = tf_r.add_paragraph()
    pr2.text = '"Pressing RESET clears the active live anomaly state and restores telemetry to 43 L/min, while preserving Ticket #104 in the historical work-order registry."'
    pr2.font.size = Pt(13)
    pr2.font.color.rgb = WHITE
    pr2.space_before = Pt(2)

    # ==========================================
    # SLIDE 9: Projected Impact
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide9)
    add_header(slide9, "Quantifying Waste to Prioritize Operational Response", "PROJECTED IMPACT")

    # Formula Box
    fbox = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.9))
    fbox.fill.solid()
    fbox.fill.fore_color.rgb = CARD_BG
    fbox.line.color.rgb = CARD_BORDER
    tf_f = fbox.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "EXCESS FLOW RATE FORMULA:  Excess (E) = max(0, Current Flow - Baseline Flow) = 94 - 43 = 51 L/min"
    pf.font.size = Pt(14)
    pf.font.bold = True
    pf.font.color.rgb = CYAN

    # 4 Volume Cards
    v_cards = [
        ("15 MIN SIMULATED", "765 L", CYAN),
        ("1 HOUR PROJECTED", "3,060 L", CYAN),
        ("6 HOURS PROJECTED", "18,360 L", AMBER),
        ("24 HOURS PROJECTED", "73,440 L", RED)
    ]
    vw = Inches(2.7)
    vgap = Inches(0.3)
    for i, (vlabel, vval, vcol) in enumerate(v_cards):
        vx = Inches(0.8) + i * (vw + vgap)
        vshape = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, vx, Inches(2.6), vw, Inches(2.0))
        vshape.fill.solid()
        vshape.fill.fore_color.rgb = CARD_BG
        vshape.line.color.rgb = vcol
        tf_v = vshape.text_frame
        tf_v.word_wrap = True
        pv1 = tf_v.paragraphs[0]
        pv1.text = vlabel
        pv1.font.size = Pt(11)
        pv1.font.bold = True
        pv1.font.color.rgb = MUTED
        pv2 = tf_v.add_paragraph()
        pv2.text = vval
        pv2.font.size = Pt(28)
        pv2.font.bold = True
        pv2.font.color.rgb = vcol
        pv2.space_before = Pt(16)

    # Financial Cost Box
    fin_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.8), Inches(11.7), Inches(2.2))
    fin_box.fill.solid()
    fin_box.fill.fore_color.rgb = CARD_BG
    fin_box.line.color.rgb = GREEN
    tf_fin = fin_box.text_frame
    tf_fin.word_wrap = True
    pfin = tf_fin.paragraphs[0]
    pfin.text = "ESTIMATED 6-HOUR FINANCIAL LOSS FORECAST: ₹826.20"
    pfin.font.size = Pt(20)
    pfin.font.bold = True
    pfin.font.color.rgb = GREEN

    pfin2 = tf_fin.add_paragraph()
    pfin2.text = "Calculation Breakdown: (18,360 L ÷ 1,000) × ₹45/kL = ₹826.20"
    pfin2.font.size = Pt(14)
    pfin2.font.color.rgb = WHITE
    pfin2.space_before = Pt(10)

    pfin3 = tf_fin.add_paragraph()
    pfin3.text = "Disclaimer: All metrics represent mathematical loss projections from simulated telemetry data."
    pfin3.font.size = Pt(11)
    pfin3.font.italic = True
    pfin3.font.color.rgb = MUTED
    pfin3.space_before = Pt(10)

    # ==========================================
    # SLIDE 10: Differentiation
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide10)
    add_header(slide10, "Not Another Dashboard. An Operational Intelligence Layer.", "DIFFERENTIATION")

    col_w = Inches(5.6)
    # Conventional
    conv_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), col_w, Inches(4.4))
    conv_box.fill.solid()
    conv_box.fill.fore_color.rgb = CARD_BG
    conv_box.line.color.rgb = CARD_BORDER
    tf_cv = conv_box.text_frame
    tf_cv.word_wrap = True
    pcv = tf_cv.paragraphs[0]
    pcv.text = "CONVENTIONAL MONITORING APPROACH"
    pcv.font.size = Pt(16)
    pcv.font.bold = True
    pcv.font.color.rgb = MUTED
    
    conv_items = [
        "Raw telemetry visualization",
        "Operator interpretation required",
        "Limited impact context provided",
        "Separate operational workflow"
    ]
    for item in conv_items:
        p = tf_cv.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(14)
        p.font.color.rgb = MUTED
        p.space_before = Pt(14)

    # AquaSentinel
    aq_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), col_w, Inches(4.4))
    aq_box.fill.solid()
    aq_box.fill.fore_color.rgb = CARD_BG
    aq_box.line.color.rgb = CYAN
    tf_aq = aq_box.text_frame
    tf_aq.word_wrap = True
    paq = tf_aq.paragraphs[0]
    paq.text = "AQUASENTINEL PROTOTYPE"
    paq.font.size = Pt(16)
    paq.font.bold = True
    paq.font.color.rgb = CYAN

    aq_items = [
        "Baseline-based anomaly detection engine",
        "AI diagnostic explanation & evidence",
        "Projected water volume & financial impact (₹)",
        "Integrated work-order ticketing workflow"
    ]
    for item in aq_items:
        p = tf_aq.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.space_before = Pt(14)

    # Positioning Banner
    pos_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    pos_box.fill.solid()
    pos_box.fill.fore_color.rgb = CARD_BG
    pos_box.line.color.rgb = CYAN
    tf_pos = pos_box.text_frame
    ppos = tf_pos.paragraphs[0]
    ppos.text = '"AquaSentinel sits on top of water telemetry as an intelligent operational layer."'
    ppos.font.size = Pt(16)
    ppos.font.bold = True
    ppos.font.color.rgb = CYAN
    ppos.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 11: Scalability & Closing
    # ==========================================
    slide11 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide11)
    add_header(slide11, "Built as a Working Prototype. Designed for Real-World Scale.", "ROADMAP & CLOSING")

    # Today vs Next
    t_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(3.6))
    t_box.fill.solid()
    t_box.fill.fore_color.rgb = CARD_BG
    t_box.line.color.rgb = CYAN
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    pt1 = tf_t.paragraphs[0]
    pt1.text = "TODAY (Working Prototype)"
    pt1.font.size = Pt(16)
    pt1.font.bold = True
    pt1.font.color.rgb = CYAN

    today_list = [
        "Simulated sensor telemetry engine",
        "1.5x baseline anomaly threshold rule",
        "AI diagnostic explanation",
        "Projected impact in ₹ INR",
        "Work-order ticketing workflow"
    ]
    for item in today_list:
        p = tf_t.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.space_before = Pt(8)

    n_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), Inches(5.6), Inches(3.6))
    n_box.fill.solid()
    n_box.fill.fore_color.rgb = CARD_BG
    n_box.line.color.rgb = CARD_BORDER
    tf_n = n_box.text_frame
    tf_n.word_wrap = True
    pn1 = tf_n.paragraphs[0]
    pn1.text = "NEXT PHASE (Real-World Deployment)"
    pn1.font.size = Pt(16)
    pn1.font.bold = True
    pn1.font.color.rgb = MUTED

    next_list = [
        "ESP32 + Pulse Flow Meters",
        "MQTT / Modbus Broker integration",
        "Multi-building facility deployment",
        "Historical real telemetry model training",
        "Automated shut-off valve triggers"
    ]
    for item in next_list:
        p = tf_n.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = MUTED
        p.space_before = Pt(8)

    # Final Pitch Banner
    pitch_box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.7))
    pitch_box.fill.solid()
    pitch_box.fill.fore_color.rgb = CARD_BG
    pitch_box.line.color.rgb = CYAN
    tf_pi = pitch_box.text_frame
    tf_pi.word_wrap = True

    ppi1 = tf_pi.paragraphs[0]
    ppi1.text = '"AquaSentinel turns water data into decisions — before waste becomes a bill."'
    ppi1.font.size = Pt(20)
    ppi1.font.bold = True
    ppi1.font.color.rgb = WHITE
    ppi1.alignment = PP_ALIGN.CENTER

    ppi2 = tf_pi.add_paragraph()
    ppi2.text = "MONITOR  →  DETECT  →  EXPLAIN  →  PREDICT  →  ACT"
    ppi2.font.size = Pt(14)
    ppi2.font.bold = True
    ppi2.font.color.rgb = CYAN
    ppi2.alignment = PP_ALIGN.CENTER
    ppi2.space_before = Pt(10)

    ppi3 = tf_pi.add_paragraph()
    ppi3.text = "Live Demo: Available during presentation  │  GitHub: github.com/yasham-06/Aquasentinal"
    ppi3.font.size = Pt(11)
    ppi3.font.color.rgb = MUTED
    ppi3.alignment = PP_ALIGN.CENTER
    ppi3.space_before = Pt(8)

    # Save output
    output_path = os.path.join(os.getcwd(), "AquaSentinel_Hackathon_Pitch.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Saved presentation to {output_path}")

if __name__ == "__main__":
    create_presentation()
